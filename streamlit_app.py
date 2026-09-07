import io, json, re, zipfile, requests
import streamlit as st
from pypdf import PdfReader
from docx import Document

st.set_page_config(page_title="Question Paper Maker",page_icon="📝",layout="wide")
st.title("📝 Question Paper Maker")
st.caption("Public web version — generate question papers, MCQ keys and proper solutions from your own source material.")

def read_file(f):
    f.seek(0)
    if f.name.lower().endswith(".pdf"):
        return "\n".join((p.extract_text() or "") for p in PdfReader(f).pages)
    if f.name.lower().endswith(".docx"):
        d=Document(f); out=[p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for r in t.rows:
                s=" | ".join(c.text.strip() for c in r.cells if c.text.strip())
                if s: out.append(s)
        return "\n".join(out)
    return f.read().decode("utf-8",errors="ignore")

def split_units(text):
    m=list(re.finditer(r"(?im)\bunit\s*[-–:]?\s*(?:[ivx]+|\d+)\b",text))
    if len(m)<5: raise ValueError("Five Unit headings were not detected. Please use Unit 1 ... Unit 5 headings.")
    out=[]
    for i,x in enumerate(m[:5]):
        end=m[i+1].start() if i+1<len(m) else len(text)
        out.append(text[x.start():end].strip())
    return out

def call_ai(prompt):
    key=st.secrets.get("GEMINI_API_KEY","")
    model=st.secrets.get("GEMINI_MODEL","gemini-2.5-flash")
    if not key: raise RuntimeError("Website owner has not configured the AI key.")
    url=f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={key}"
    payload={"contents":[{"parts":[{"text":prompt}]}],
             "generationConfig":{"temperature":0.25,"responseMimeType":"application/json"}}
    r=requests.post(url,json=payload,timeout=240)
    if not r.ok: raise RuntimeError(f"AI service error {r.status_code}: {r.text[:300]}")
    return r.json()["candidates"][0]["content"]["parts"][0]["text"]

def get_json(s):
    s=re.sub(r"^```(?:json)?\s*","",s.strip()); s=re.sub(r"\s*```$","",s)
    return json.loads(s)

def generate(unit_no,unit,reference,co,difficulty,used):
    # Keep prompt size practical for a public MVP.
    ref=reference[:30000]
    prompt=f"""Act as a university examination question-paper setter.
Use ONLY the supplied Unit syllabus and reference material. Do not add outside topics.
Avoid questions already used. Difficulty: {difficulty}.
UNIT {unit_no} SYLLABUS:
{unit}
REFERENCE MATERIAL:
{ref}
COURSE OUTCOMES:
{co}
ALREADY USED:
{used[-7000:]}

Generate exactly 10 MCQs and exactly 2 alternative descriptive main questions.
Each descriptive main question is 5 marks split into part a = 3 and part b = 2.
MCQs need exactly four plausible options and one correct answer.
Solutions must be actual model answers, never phrases like "award marks for".
Return only valid JSON:
{{"mcqs":[{{"question":"","options":{{"A":"","B":"","C":"","D":""}},"answer":"A","bt":"L2","co":"CO1"}}],
"descriptive":[{{"a":{{"question":"","marks":3,"solution":""}},"b":{{"question":"","marks":2,"solution":""}},"bt":"L2","co":"CO1"}},
{{"a":{{"question":"","marks":3,"solution":""}},"b":{{"question":"","marks":2,"solution":""}},"bt":"L3","co":"CO1"}}]}}"""
    d=get_json(call_ai(prompt))
    if len(d.get("mcqs",[]))!=10 or len(d.get("descriptive",[]))!=2:
        raise ValueError("AI returned an incomplete unit. Please regenerate.")
    return d

def balance(qs):
    target=(["A","B","C","D"]*13)[:len(qs)]
    for q,d in zip(qs,target):
        c=q["answer"].upper()
        if c!=d:
            q["options"][c],q["options"][d]=q["options"][d],q["options"][c]; q["answer"]=d
    return qs

def make_key_doc(paper,title,code,setno):
    d=Document(); d.add_heading("MCQ Answer Key",0)
    d.add_paragraph(f"Course: {code} – {title}     Set: {setno}")
    t=d.add_table(rows=1,cols=5); t.style="Table Grid"
    for c,x in zip(t.rows[0].cells,["Q.No.","Answer","Unit","BT","CO"]): c.text=x
    for i,q in enumerate(paper["mcqs"],1):
        r=t.add_row().cells
        for c,x in zip(r,[i,q["answer"],q["unit"],q.get("bt",""),q.get("co","")]): c.text=str(x)
    o=io.BytesIO(); d.save(o); return o.getvalue()

def make_qp_doc(paper,title,code,setno):
    d=Document(); d.add_heading("Question Paper",0)
    d.add_paragraph(f"Course: {code} – {title}     Set: {setno}")
    d.add_heading("Part A – MCQs",1)
    for i,q in enumerate(paper["mcqs"],1):
        d.add_paragraph(f"Q{i}. {q['question']}")
        d.add_paragraph(f"A) {q['options']['A']}     B) {q['options']['B']}\nC) {q['options']['C']}     D) {q['options']['D']}")
    d.add_heading("Part B – Descriptive Questions",1)
    n=1
    for u,pair in enumerate(paper["descriptive"],1):
        d.add_paragraph(f"UNIT {u}",style="Heading 2")
        for j,q in enumerate(pair):
            d.add_paragraph(f"Q{n}. a) {q['a']['question']} ({q['a']['marks']} marks)")
            d.add_paragraph(f"     b) {q['b']['question']} ({q['b']['marks']} marks)")
            if j==0: d.add_paragraph("OR")
            n+=1
    o=io.BytesIO(); d.save(o); return o.getvalue()

def make_sol_doc(paper,title,code,setno):
    d=Document(); d.add_heading("Scheme & Solution",0)
    d.add_paragraph(f"Course: {code} – {title}     Set: {setno}")
    t=d.add_table(rows=1,cols=3); t.style="Table Grid"
    for c,x in zip(t.rows[0].cells,["Q.No.","Solution","Marks"]): c.text=x
    n=1
    for pair in paper["descriptive"]:
        for q in pair:
            r=t.add_row().cells; r[0].text=str(n)
            r[1].text=f"a) {q['a']['solution']}\n\nb) {q['b']['solution']}"
            r[2].text="5"; n+=1
    o=io.BytesIO(); d.save(o); return o.getvalue()


def get_bytes(f):
    f.seek(0)
    return f.read()

def set_cell(c, text):
    c.text = str(text)

def fill_qp_template(template_bytes, paper, meta):
    d=Document(io.BytesIO(template_bytes))
    if len(d.tables)<4:
        raise ValueError("Question-paper template structure is not compatible with the current Nitte SEE format.")
    t0,t1,t2,t3=d.tables[:4]
    if len(t0.rows)>=2:
        set_cell(t0.rows[0].cells[1],meta.get("setter","")); set_cell(t0.rows[0].cells[3],meta.get("set",""))
        set_cell(t0.rows[1].cells[1],meta.get("department","")); set_cell(t0.rows[1].cells[3],meta.get("date",""))
    if len(t1.rows)>=5:
        set_cell(t1.rows[2].cells[0],f'{meta.get("semester","")} Semester B.Sc (CBCS) Degree Examinations')
        set_cell(t1.rows[3].cells[0],f'Academic Year: {meta.get("year","")}')
        set_cell(t1.rows[4].cells[0],f'{meta.get("code","")} – {meta.get("title","")} ({meta.get("scheme","")})\n(For {meta.get("branches","")})')
    # MCQs: current SEE template uses 3 rows/question
    for i,q in enumerate(paper["mcqs"][:50]):
        r=i*3
        if r+2>=len(t2.rows): break
        set_cell(t2.rows[r].cells[0],f"{i+1}."); set_cell(t2.rows[r].cells[1],q["question"])
        set_cell(t2.rows[r+1].cells[0],""); set_cell(t2.rows[r+1].cells[1],"A)"); set_cell(t2.rows[r+1].cells[2],q["options"]["A"]); set_cell(t2.rows[r+1].cells[3],"B)"); set_cell(t2.rows[r+1].cells[4],q["options"]["B"])
        set_cell(t2.rows[r+2].cells[0],""); set_cell(t2.rows[r+2].cells[1],"C)"); set_cell(t2.rows[r+2].cells[2],q["options"]["C"]); set_cell(t2.rows[r+2].cells[3],"D)"); set_cell(t2.rows[r+2].cells[4],q["options"]["D"])
    rows=[((1,2),(4,5)),((7,8),(10,11)),((13,14),(16,17)),((19,20),(22,23)),((25,26),(28,29))]
    qn=1
    for ui,pair in enumerate(paper["descriptive"][:5]):
        for alt,(ra,rb) in enumerate(rows[ui]):
            q=pair[alt]
            for rr,part,label in [(ra,q["a"],"a)"),(rb,q["b"],"b)")]:
                set_cell(t3.rows[rr].cells[0],str(qn) if label=="a)" else "")
                set_cell(t3.rows[rr].cells[1],label); set_cell(t3.rows[rr].cells[2],part["question"]); set_cell(t3.rows[rr].cells[3],part["marks"])
                set_cell(t3.rows[rr].cells[4],q.get("bt","")); set_cell(t3.rows[rr].cells[5],q.get("co","")); set_cell(t3.rows[rr].cells[6],q.get("po",""))
            qn+=1
    o=io.BytesIO(); d.save(o); return o.getvalue()

def fill_key_template(template_bytes,paper,meta):
    d=Document(io.BytesIO(template_bytes))
    if not d.tables: raise ValueError("MCQ answer-key template has no table.")
    for p in d.paragraphs:
        if "Program:" in p.text and "Semester:" in p.text: p.text=f'Program: B.Sc    Semester: {meta.get("semester","")}    QP SET No.: {meta.get("set","")}'
        elif "Course Code:" in p.text and "Course Title:" in p.text: p.text=f'Course Code: {meta.get("code","")}    Course Title: {meta.get("title","")}'
        elif p.text.strip().startswith("Name of the Faculty:"): p.text=f'Name of the Faculty: {meta.get("setter","")}'
        elif p.text.strip().startswith("Affiliation:"): p.text=f'Affiliation: {meta.get("department","")}'
    t=d.tables[0]; ans=[q["answer"].upper() for q in paper["mcqs"][:50]]
    for r in range(1,min(11,len(t.rows))):
        for b in range(5):
            n=r+b*10
            if n<=len(ans) and b*2+1<len(t.rows[r].cells):
                set_cell(t.rows[r].cells[b*2],f"{n}."); set_cell(t.rows[r].cells[b*2+1],ans[n-1])
    o=io.BytesIO(); d.save(o); return o.getvalue()

def fill_sol_template(template_bytes,paper,meta):
    d=Document(io.BytesIO(template_bytes))
    if len(d.tables)<2: raise ValueError("Scheme & Solution template structure is not compatible.")
    h=d.tables[0]
    if len(h.rows)>=2:
        set_cell(h.rows[0].cells[1],meta.get("title","")); set_cell(h.rows[0].cells[3],meta.get("code",""))
        set_cell(h.rows[1].cells[1],meta.get("setter","")); set_cell(h.rows[1].cells[3],meta.get("set",""))
    t=d.tables[1]
    while len(t.rows)>1: t._tbl.remove(t.rows[-1]._tr)
    n=1
    for pair in paper["descriptive"]:
        for q in pair:
            r=t.add_row().cells; set_cell(r[0],n); set_cell(r[1],f'a) {q["a"]["solution"]}\n\nb) {q["b"]["solution"]}'); set_cell(r[2],5); n+=1
    o=io.BytesIO(); d.save(o); return o.getvalue()

with st.sidebar:
    sets=st.slider("Number of sets",1,4,4)
    difficulty=st.selectbox("Difficulty",["Easy","Easy–Moderate","Moderate","Moderate–Hard"],1)

t1,t2,t3,t4=st.tabs(["1. Course Setup","2. Sources & Templates","3. Generate & Review","4. Export"])
with t1:
    a,b,c=st.columns(3); title=a.text_input("Course title"); code=b.text_input("Course code"); semester=c.text_input("Semester")
    a,b,c=st.columns(3); year=a.text_input("Academic year"); scheme=b.text_input("Scheme",value="2025 Scheme"); branches=c.text_input("Branches",value="CAPT / CAFD")
    a,b,c=st.columns(3); setter=a.text_input("QP setter / Faculty"); department=b.text_input("Department / Affiliation"); date=c.text_input("Date")
    co=st.text_area("Course Outcomes",height=140,placeholder="CO1 ...\nCO2 ...")

with t2:
    st.subheader("Source material")
    syllabus=st.file_uploader("Syllabus",["pdf","docx","txt"],key="sy")
    reference=st.file_uploader("Reference book / notes",["pdf","docx","txt"],key="ref")
    st.subheader("University Word templates")
    qp_template=st.file_uploader("Question Paper Template (.docx)",["docx"],key="qpt")
    key_template=st.file_uploader("MCQ Answer Key Template (.docx)",["docx"],key="akt")
    sol_template=st.file_uploader("Scheme & Solution Template (.docx)",["docx"],key="sot")
    if syllabus: st.session_state["syllabus"]=read_file(syllabus); st.success("Syllabus loaded.")
    if reference: st.session_state["reference"]=read_file(reference); st.success("Reference material loaded.")

with t3:
    if st.button("Generate Question Papers",type="primary",use_container_width=True):
        if not st.session_state.get("syllabus") or not st.session_state.get("reference") or not co.strip():
            st.error("Enter COs and upload syllabus + reference material.")
        else:
            try:
                units=split_units(st.session_state["syllabus"]); papers=[]; used=""; bar=st.progress(0)
                for sidx in range(sets):
                    mc=[]; desc=[]
                    for i,u in enumerate(units,1):
                        x=generate(i,u,st.session_state["reference"],co,difficulty,used)
                        for q in x["mcqs"]:
                            q["unit"]=i; q.setdefault("po",""); mc.append(q); used+="\n"+q["question"]
                        desc.append(x["descriptive"])
                        for q in x["descriptive"]:
                            q.setdefault("po",""); used+="\n"+q["a"]["question"]+"\n"+q["b"]["question"]
                    papers.append({"mcqs":balance(mc),"descriptive":desc}); bar.progress((sidx+1)/sets)
                st.session_state["papers"]=papers; st.success("Generation complete.")
            except Exception as e: st.error(str(e))
    if st.session_state.get("papers"):
        n=st.selectbox("Review Set",range(1,len(st.session_state["papers"])+1)); p=st.session_state["papers"][n-1]
        st.markdown("### MCQs")
        for i,q in enumerate(p["mcqs"],1):
            with st.expander(f'Q{i}. {q["question"]}'):
                q["question"]=st.text_area("Question",q["question"],key=f"q{n}_{i}")
                cols=st.columns(4)
                for j,L in enumerate("ABCD"): q["options"][L]=cols[j].text_input(L,q["options"][L],key=f"o{n}_{i}_{L}")
                q["answer"]=st.selectbox("Correct answer",list("ABCD"),index=list("ABCD").index(q["answer"]),key=f"a{n}_{i}")
        st.markdown("### Descriptive questions and solutions")
        qn=1
        for u,pair in enumerate(p["descriptive"],1):
            for alt,q in enumerate(pair,1):
                with st.expander(f"Question {qn} – Unit {u}"):
                    q["a"]["question"]=st.text_area("a) Question",q["a"]["question"],key=f"daq{n}_{u}_{alt}")
                    q["a"]["solution"]=st.text_area("a) Solution",q["a"]["solution"],height=110,key=f"das{n}_{u}_{alt}")
                    q["b"]["question"]=st.text_area("b) Question",q["b"]["question"],key=f"dbq{n}_{u}_{alt}")
                    q["b"]["solution"]=st.text_area("b) Solution",q["b"]["solution"],height=110,key=f"dbs{n}_{u}_{alt}")
                qn+=1

with t4:
    ps=st.session_state.get("papers",[])
    if not ps: st.info("Generate papers first.")
    elif not all([qp_template,key_template,sol_template]): st.info("Upload all three Word templates in Sources & Templates.")
    else:
        if st.button("Build DOCX Package Using Uploaded Templates",type="primary",use_container_width=True):
            try:
                qb,kb,sb=get_bytes(qp_template),get_bytes(key_template),get_bytes(sol_template)
                out=io.BytesIO()
                with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as z:
                    for i,p in enumerate(ps,1):
                        meta={"title":title,"code":code,"semester":semester,"year":year,"scheme":scheme,"branches":branches,"setter":setter,"department":department,"date":date,"set":str(i)}
                        z.writestr(f"Set_{i}/Question_Paper_Set_{i}.docx",fill_qp_template(qb,p,meta))
                        z.writestr(f"Set_{i}/MCQ_Answer_Key_Set_{i}.docx",fill_key_template(kb,p,meta))
                        z.writestr(f"Set_{i}/Scheme_Solution_Set_{i}.docx",fill_sol_template(sb,p,meta))
                st.session_state["zip"]=out.getvalue(); st.success("Template-based package ready.")
            except Exception as e: st.error(f"Export failed: {e}")
        if st.session_state.get("zip"):
            st.download_button("Download All Sets",st.session_state["zip"],"Question_Paper_Package.zip","application/zip",use_container_width=True)

st.divider()
st.caption("Public Web v2.1 • Restores template uploads and template-based DOCX export from the local app.")
